#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
قالب Excel للفواتير - نظام التميز للهاتف النقال
إنشاء وقراءة قوالب Excel للفواتير
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import os
from datetime import datetime

def create_invoice_template():
    """إنشاء قالب Excel للفواتير"""
    
    # إنشاء ملف Excel جديد
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "قالب الفواتير"
    
    # إعداد الأنماط
    header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                   top=Side(style='thin'), bottom=Side(style='thin'))
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    # معلومات الشركة
    ws['A1'] = 'التميز للهاتف النقال وقطع الغيار'
    ws['A1'].font = Font(name='Arial', size=16, bold=True)
    ws['A1'].alignment = center_alignment
    ws.merge_cells('A1:H1')
    
    ws['A2'] = 'Excellence Mobile & Spare Parts'
    ws['A2'].font = Font(name='Arial', size=12, bold=True)
    ws['A2'].alignment = center_alignment
    ws.merge_cells('A2:H2')
    
    # معلومات الفاتورة
    ws['A4'] = 'رقم الفاتورة:'
    ws['B4'] = '[INV-NUMBER]'
    ws['E4'] = 'التاريخ:'
    ws['F4'] = '[DATE]'
    
    ws['A5'] = 'العميل:'
    ws['B5'] = '[CUSTOMER-NAME]'
    ws['E5'] = 'الهاتف:'
    ws['F5'] = '[CUSTOMER-PHONE]'
    
    # رأس جدول المنتجات
    headers = ['م', 'اسم المنتج', 'رمز المنتج', 'الكمية', 'سعر الوحدة', 'المجموع', 'ملاحظات', 'حالة المردود']
    row = 7
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = center_alignment
    
    # صفوف المنتجات (30 صف)
    for i in range(1, 31):
        row_num = row + i
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.border = border
            if col == 1:  # رقم الصف
                cell.value = i
                cell.alignment = center_alignment
            elif col == 8:  # حالة المردود
                cell.value = 'لا'  # القيمة الافتراضية
                cell.alignment = center_alignment
    
    # إجماليات الفاتورة
    total_row = row + 32
    ws[f'E{total_row}'] = 'المجموع الفرعي:'
    ws[f'F{total_row}'] = '[SUBTOTAL]'
    
    ws[f'E{total_row + 1}'] = 'الخصم:'
    ws[f'F{total_row + 1}'] = '[DISCOUNT]'
    
    ws[f'E{total_row + 2}'] = 'المردودات:'
    ws[f'F{total_row + 2}'] = '[RETURNS]'
    
    ws[f'E{total_row + 3}'] = 'الضريبة (15%):'
    ws[f'F{total_row + 3}'] = '[TAX]'
    
    ws[f'E{total_row + 4}'] = 'الإجمالي النهائي:'
    ws[f'F{total_row + 4}'] = '[TOTAL]'
    ws[f'F{total_row + 4}'].font = Font(bold=True)
    
    # تعديل عرض الأعمدة
    column_widths = [5, 25, 15, 8, 12, 12, 20, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    
    # إضافة ورقة التعليمات
    instructions_sheet = wb.create_sheet("تعليمات الاستخدام")
    
    instructions = [
        "تعليمات استخدام قالب الفواتير:",
        "",
        "1. املأ بيانات المنتجات في الصفوف المخصصة",
        "2. الأعمدة المطلوبة:",
        "   - اسم المنتج: اسم المنتج كما هو في النظام",
        "   - رمز المنتج: SKU أو الباركود",
        "   - الكمية: الكمية المباعة",
        "   - سعر الوحدة: سعر القطعة الواحدة",
        "   - المجموع: سيتم حسابه تلقائياً (الكمية × السعر)",
        "   - حالة المردود: نعم/لا (للمنتجات المردودة)",
        "",
        "3. معلومات إضافية:",
        "   - رقم الفاتورة: سيتم إنشاؤه تلقائياً",
        "   - التاريخ: تاريخ اليوم افتراضياً",
        "   - معلومات العميل: اختيارية",
        "",
        "4. ملاحظات مهمة:",
        "   - تأكد من صحة أسماء المنتجات",
        "   - تأكد من وجود المنتجات في النظام",
        "   - الأسعار بالريال السعودي",
        "   - المردودات ستؤثر على المخزون",
        "",
        "5. بعد الانتهاء:",
        "   - احفظ الملف",
        "   - ارفعه للنظام عبر صفحة المشتريات",
        "   - سيتم معالجة الفاتورة تلقائياً"
    ]
    
    for i, instruction in enumerate(instructions, 1):
        instructions_sheet[f'A{i}'] = instruction
        if i == 1:
            instructions_sheet[f'A{i}'].font = Font(bold=True, size=14)
        elif instruction.startswith(('1.', '2.', '3.', '4.', '5.')):
            instructions_sheet[f'A{i}'].font = Font(bold=True)
    
    # تعديل عرض العمود في ورقة التعليمات
    instructions_sheet.column_dimensions['A'].width = 60
    
    # حفظ الملف
    template_path = 'invoice_template.xlsx'
    wb.save(template_path)
    
    return template_path

def parse_excel_invoice(file_path):
    """قراءة وتحليل فاتورة من ملف Excel"""
    try:
        # قراءة الملف
        df = pd.read_excel(file_path, sheet_name=0, header=6)  # البيانات تبدأ من الصف 7
        
        # تنظيف البيانات
        df = df.dropna(subset=['اسم المنتج'])  # إزالة الصفوف الفارغة
        df = df[df['اسم المنتج'].str.strip() != '']  # إزالة الصفوف الفارغة
        
        # قراءة معلومات الفاتورة من الخلايا العلوية
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        invoice_info = {
            'invoice_number': str(ws['B4'].value) if ws['B4'].value else None,
            'date': ws['F4'].value if ws['F4'].value else datetime.now().isoformat(),
            'customer_name': str(ws['B5'].value) if ws['B5'].value else 'عميل عادي',
            'customer_phone': str(ws['F5'].value) if ws['F5'].value else ''
        }
        
        # تحويل البيانات إلى قائمة منتجات
        products = []
        for index, row in df.iterrows():
            if pd.isna(row['اسم المنتج']) or str(row['اسم المنتج']).strip() == '':
                continue
                
            # التحقق من المردود
            is_returned = str(row.get('حالة المردود', 'لا')).strip().lower() in ['نعم', 'yes', '1', 'true']
            
            product = {
                'name': str(row['اسم المنتج']).strip(),
                'sku': str(row['رمز المنتج']).strip() if pd.notna(row['رمز المنتج']) else '',
                'quantity': int(row['الكمية']) if pd.notna(row['الكمية']) else 1,
                'unit_price': float(row['سعر الوحدة']) if pd.notna(row['سعر الوحدة']) else 0,
                'total': float(row['المجموع']) if pd.notna(row['المجموع']) else 0,
                'notes': str(row['ملاحظات']).strip() if pd.notna(row['ملاحظات']) else '',
                'is_returned': is_returned,
                'return_quantity': int(row['الكمية']) if is_returned else 0
            }
            
            # حساب المجموع إذا لم يكن محدداً
            if product['total'] == 0:
                product['total'] = product['quantity'] * product['unit_price']
            
            products.append(product)
        
        return {
            'success': True,
            'invoice_info': invoice_info,
            'products': products,
            'products_count': len(products)
        }
        
    except Exception as e:
        return {
            'success': False,
            'error': str(e),
            'products': []
        }

def create_sample_invoice():
    """إنشاء فاتورة عينة للاختبار"""
    template_path = create_invoice_template()
    
    # تحميل القالب
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    
    # ملء بيانات عينة
    ws['B4'] = 'INV-20250105-0001'
    ws['F4'] = datetime.now().strftime('%Y-%m-%d')
    ws['B5'] = 'أحمد محمد'
    ws['F5'] = '0501234567'
    
    # إضافة منتجات عينة
    sample_products = [
        ['iPhone 15 Pro Max', 'IP15PM256', 2, 4200, 8400, 'أسود 256GB', 'لا'],
        ['سماعة AirPods Pro', 'APP3GEN', 1, 899, 899, 'الجيل الثالث', 'لا'],
        ['حافظة iPhone 15', 'CASE15PM', 3, 45, 135, 'شفافة', 'نعم'],  # مردود
        ['شاحن سريع 20W', 'CHRG20W', 2, 85, 170, 'أبيض أصلي', 'لا']
    ]
    
    start_row = 8
    for i, product in enumerate(sample_products):
        row = start_row + i
        ws[f'A{row}'] = i + 1
        ws[f'B{row}'] = product[0]  # اسم المنتج
        ws[f'C{row}'] = product[1]  # رمز المنتج
        ws[f'D{row}'] = product[2]  # الكمية
        ws[f'E{row}'] = product[3]  # سعر الوحدة
        ws[f'F{row}'] = product[4]  # المجموع
        ws[f'G{row}'] = product[5]  # ملاحظات
        ws[f'H{row}'] = product[6]  # حالة المردود
    
    # حساب الإجماليات
    subtotal = 8400 + 899 + 135 + 170
    returns = 135  # حافظة iPhone مردودة
    net_subtotal = subtotal - returns
    tax = net_subtotal * 0.15
    total = net_subtotal + tax
    
    # ملء الإجماليات
    total_row = 39
    ws[f'F{total_row}'] = subtotal
    ws[f'F{total_row + 1}'] = 0  # خصم
    ws[f'F{total_row + 2}'] = returns
    ws[f'F{total_row + 3}'] = round(tax, 2)
    ws[f'F{total_row + 4}'] = round(total, 2)
    
    # حفظ الفاتورة العينة
    sample_path = 'sample_invoice.xlsx'
    wb.save(sample_path)
    
    return sample_path

if __name__ == "__main__":
    print("إنشاء قالب Excel للفواتير...")
    template_path = create_invoice_template()
    print(f"تم إنشاء القالب: {template_path}")
    
    print("إنشاء فاتورة عينة...")
    sample_path = create_sample_invoice()
    print(f"تم إنشاء الفاتورة العينة: {sample_path}")
    
    print("اختبار قراءة الفاتورة العينة...")
    result = parse_excel_invoice(sample_path)
    if result['success']:
        print(f"تم قراءة {result['products_count']} منتج بنجاح")
        print("المنتجات:")
        for product in result['products']:
            status = "مردود" if product['is_returned'] else "عادي"
            print(f"- {product['name']}: {product['quantity']} × {product['unit_price']} = {product['total']} ({status})")
    else:
        print(f"خطأ في قراءة الفاتورة: {result['error']}")